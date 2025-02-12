import tkinter as tk
from tkinter import messagebox, filedialog
from datetime import datetime
import json
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from tkinter import ttk


class DepoYonetimi:
    def __init__(self, root):
        self.depo = {}
        self.urun_istekleri = []  # Ürün isteklerini saklayacak liste
        self.root = root
        self.root.title("Depo Yönetim Sistemi")

        # Ekranın ortasında kare pencere
        window_width = 600
        window_height = 600
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        position_top = int(screen_height / 2 - window_height / 2)
        position_right = int(screen_width / 2 - window_width / 2)
        self.root.geometry(f'{window_width}x{window_height}+{position_right}+{position_top}')
        
        self.root.configure(bg='#f0f0f0')

        self.dosyayi_yukle()
        self.ana_menu()

    def show_author_label(self, parent):
        label = tk.Label(parent, text="Created by Ertuğrul Yıldırım for MSS", bg='#f0f0f0', fg='black', font=('Arial', 7))
        label.pack(side=tk.RIGHT, anchor='e')


    def ana_menu(self):
        for widget in self.root.winfo_children():
            widget.destroy()

        style = ttk.Style()
        style.theme_use('clam')
        style.configure('TButton', background='#66ff66', foreground='black', font=('Arial', 10), padding=10)
        
        self.sifirla_button = tk.Button(self.root, text="Depoyu Sıfırla", command=self.depo_sifirla_uyari, bg='#ff6666', fg='white', font=('Arial', 8))
        self.sifirla_button.place(x=10, y=10)

        self.destek_button = tk.Button(self.root, text="Destek", command=self.destek_penceresi, bg='#ff6666', fg='white', font=('Arial', 8))
        self.destek_button.place(x=10, y=50)

        self.aylik_ihtiyac_button = ttk.Button(self.root, text="Aylık İhtiyaç Listesi", command=self.ac_aylik_ihtiyac, style='TButton')
        self.aylik_ihtiyac_button.place(x=450, y=10)

        frame = tk.Frame(self.root, bg='#f0f0f0')
        frame.pack(pady=20, padx=20)

        #self.ekle_button = ttk.Button(frame, text="Ürün Ekle", command=self.urun_ekleme_penceresi, style='TButton')
        #self.ekle_button.grid(row=0, column=0, padx=10, pady=10)

        self.cikar_button = ttk.Button(frame, text="Ürün Çıkar", command=self.urun_cikarma_penceresi, style='TButton')
        self.cikar_button.grid(row=1, column=0, padx=10, pady=10)

        self.export_button = ttk.Button(frame, text="İşlemleri Excel'e Aktar", command=self.excel_aktar, style='TButton')
        self.export_button.grid(row=2, column=0, padx=10, pady=10)

        self.urun_iste_button = ttk.Button(frame, text="Ürün İste", command=self.urun_isteme_penceresi, style='TButton')
        self.urun_iste_button.grid(row=3, column=0, padx=10, pady=10)

        self.stok_listbox = tk.Listbox(self.root, width=50, height=20, bg='#e6e6fa', fg='black')
        self.stok_listbox.pack(side=tk.LEFT, padx=10, pady=10)

        self.stok_button = ttk.Button(self.root, text="Stok Durumunu Göster", command=self.stok_goster, style='TButton')
        style.configure('TButton', background='#d3d3d3', font=('Arial', 8))
        self.stok_button.place(x=450, y=200)

        self.show_author_label(self.root)
        self.stok_goster()

    def destek_penceresi(self):
        self.pencere = tk.Toplevel(self.root)
        self.pencere.title("Destek Penceresi")
        self.pencere.geometry("400x300")
        self.pencere.configure(bg='#f0f0f0')

        destek_mesaj = tk.Text(self.pencere, wrap=tk.WORD, bg='#f0f0f0', fg='black', font=('Arial', 12))
        destek_mesaj.pack(expand=True, fill=tk.BOTH, pady=10, padx=10)
        destek_mesaj.insert(tk.END, "Bu program M.Ertuğrul Yıldırım tarafından 24.01.2025 Tarihinde Murat Ticaret Çatısı altında olan MSS Savunma için geliştirilmiştir. M.Ertuğrul Yıldırım Elektrik Elektronik Mühendisi ve Bilgisayar porgramcısıdır. Bu tarihte MSS Savunma Sakarya Fabrikada Malzeme Planlama Mühendisi olarak çalışmaktadır. Destek için ertugrull_yldrm@outlook.com e-posta adresi ile iletişime geçebilirsiniz.")
        destek_mesaj.configure(state='disabled')

    def ac_aylik_ihtiyac(self):
        pencere = tk.Toplevel(self.root)
        pencere.title("Aylık İhtiyaç Listesi")
        pencere.geometry("800x600")

        # Aylık İhtiyaç Listesi kodunu burada çalıştırıyoruz.
        frame = tk.Frame(pencere)
        frame.pack(pady=20)

        self.label_file1 = tk.Label(frame, text="Müşteri İhtiyaçları seçilmedi")
        self.label_file1.grid(row=0, column=1)

        self.label_file2 = tk.Label(frame, text="SCADAM seçilmedi")
        self.label_file2.grid(row=1, column=1)

        btn_file1 = tk.Button(frame, text="Müşteri İhtiyaçları", command=self.select_file1)
        btn_file1.grid(row=0, column=0, padx=10)

        btn_file2 = tk.Button(frame, text="SCADAM", command=self.select_file2)
        btn_file2.grid(row=1, column=0, padx=10)

        btn_compare = tk.Button(frame, text="Karşılaştır", command=self.compare_files)
        btn_compare.grid(row=2, columnspan=2, pady=20)

        self.label_result = tk.Label(pencere, text="")
        self.label_result.pack()


    def select_file1(self):
        file_path = filedialog.askopenfilename()
        self.label_file1.config(text=file_path)
        return file_path

    def select_file2(self):
        file_path = filedialog.askopenfilename()
        self.label_file2.config(text=file_path)
        return file_path

    def save_file(self):
        file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        return file_path

    def compare_files(self):
        file1_path = self.label_file1.cget("text")
        file2_path = self.label_file2.cget("text")
        save_path = self.save_file()

        try:
            df1 = pd.read_excel(file1_path, header=1)
            df2 = pd.read_excel(file2_path)
        except Exception as e:
            self.label_result.config(text=f"Dosyaları okurken hata oluştu: {e}")
            return

        sixth_column = df1.iloc[:, 5]
        thirty_third_column = df1.iloc[:, 32]
        thirty_third_column = pd.to_numeric(thirty_third_column, errors='coerce')

        result_data = []

        for i, value in thirty_third_column.items():
            if pd.isna(value):
                continue

            if value < 0:
                f_value = sixth_column.iloc[i]
                new_row = {'Hedef Ambar': "", 'Ekip Kullanıcı Adı': "", 'Malzeme Tanımı': f_value, 'İstenecek Miktar': "", 'Eksik': value, 'Hedef Ambardaki Stok': ""}

                matching_row = df2[df2.iloc[:, 2] == f_value]
                if not matching_row.empty:
                    new_row['Hedef Ambardaki Stok'] = matching_row.iloc[0, 6]
                    new_row['Hedef Ambar'] = matching_row.iloc[0, 5]

                result_data.append(new_row)

        result_df = pd.DataFrame(result_data)

        writer = pd.ExcelWriter(save_path, engine='openpyxl')
        result_df.to_excel(writer, index=False, sheet_name='Sonuçlar')
        workbook  = writer.book
        worksheet = writer.sheets['Sonuçlar']

        header_fill = PatternFill(start_color='0000FF', end_color='0000FF', fill_type='solid')
        for cell in worksheet[1]:
            cell.fill = header_fill

        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal='center', vertical='center')

        writer.close()

        self.label_result.config(text=f"Karşılaştırma tamamlandı ve sonuçlar '{save_path}' dosyasına kaydedildi.")

    def show_author_label(self, parent):
        label = tk.Label(parent, text="Created by Ertuğrul Yıldırım for MSS", bg='#f0f0f0', fg='black', font=('Arial', 7))
        label.pack(side=tk.RIGHT, fill=tk.X)

    def urun_isteme_penceresi(self):

        self.pencere = tk.Toplevel(self.root)
        self.pencere.title("Ürün İste")
        self.pencere.geometry("600x600")
        self.pencere.configure(bg='#f0f0f0')

        self.ambar_label = tk.Label(self.pencere, text="İstenecek Ambar:", bg='#f0f0f0')
        self.ambar_label.grid(row=0, column=0, padx=10, pady=5, sticky='w')
        self.ambar_var = tk.StringVar()
        self.ambar_secimleri = ["0100", "1600", "1602", "1603", "3400", "3500", "5400", "5800", "7400", "MAIN"]
        for idx, secim in enumerate(self.ambar_secimleri):
            button = tk.Button(self.pencere, text=secim, command=lambda s=secim: self.ambar_sec(s), bg='#f0f0f0', fg='black')
            button.grid(row=0, column=1+idx, padx=5, pady=5)

        self.kullanici_adi_label = tk.Label(self.pencere, text="EKİP KULLANICI ADI:", bg='#f0f0f0')
        self.kullanici_adi_label.grid(row=1, column=0, padx=10, pady=5, sticky='w')
        self.kullanici_adi_entry = tk.Entry(self.pencere)
        self.kullanici_adi_entry.grid(row=1, column=1, columnspan=5, padx=10, pady=5)

        self.mrp_label = tk.Label(self.pencere, text="MRP:", bg='#f0f0f0')
        self.mrp_label.grid(row=2, column=0, padx=10, pady=5, sticky='w')
        self.mrp_entry = tk.Entry(self.pencere)
        self.mrp_entry.grid(row=2, column=1, columnspan=5, padx=10, pady=5)

        self.miktar_label = tk.Label(self.pencere, text="MİKTAR:", bg='#f0f0f0')
        self.miktar_label.grid(row=3, column=0, padx=10, pady=5, sticky='w')
        self.miktar_entry = tk.Entry(self.pencere)
        self.miktar_entry.grid(row=3, column=1, columnspan=5, padx=10, pady=5)

        self.talep_label = tk.Label(self.pencere, text="Talep Edildi Mi?", bg='#f0f0f0')
        self.talep_label.grid(row=4, column=0, padx=10, pady=5, sticky='w')
        self.talep_var = tk.StringVar()
        for idx, secim in enumerate(["Evet", "Hayır"]):
            button = tk.Button(self.pencere, text=secim, command=lambda s=secim: self.talep_sec(s), bg='#f0f0f0', fg='black')
            button.grid(row=4, column=1+idx, padx=5, pady=5)

        self.tarih_label = tk.Label(self.pencere, text="TALEP EDİLEN TARİH:", bg='#f0f0f0')
        self.tarih_label.grid(row=5, column=0, padx=10, pady=5, sticky='w')
        self.tarih_entry = tk.Entry(self.pencere)
        self.tarih_entry.grid(row=5, column=1, columnspan=5, padx=10, pady=5)

        self.sevk_sekli_label = tk.Label(self.pencere, text="SEVK ŞEKLİ:", bg='#f0f0f0')
        self.sevk_sekli_label.grid(row=6, column=0, padx=10, pady=5, sticky='w')
        self.sevk_sekli_var = tk.StringVar()
        for idx, secim in enumerate(["Ring", "Kargo", "Otobüs"]):
            button = tk.Button(self.pencere, text=secim, command=lambda s=secim: self.sevk_sec(s), bg='#f0f0f0', fg='black')
            button.grid(row=6, column=1+idx, padx=5, pady=5)

        self.yerli_ithal_label = tk.Label(self.pencere, text="Yerli mi İthal mi?", bg='#f0f0f0')
        self.yerli_ithal_label.grid(row=7, column=0, padx=10, pady=5, sticky='w')
        self.yerli_ithal_var = tk.StringVar()
        for idx, secim in enumerate(["Yerli", "İthal"]):
            button = tk.Button(self.pencere, text=secim, command=lambda s=secim: self.yerli_ithal_sec(s), bg='#f0f0f0', fg='black')
            button.grid(row=7, column=1+idx, padx=5, pady=5)

        self.talep_button = ttk.Button(self.pencere, text="Talep Gönder", command=self.urun_iste, style='TButton')
        self.talep_button.grid(row=8, column=0, columnspan=6, pady=20)

        # Talep listesi için listbox ve Excel'e aktar butonu
        self.talep_listbox = tk.Listbox(self.pencere, width=50, height=10, bg='#e6e6fa', fg='black')
        self.talep_listbox.grid(row=9, column=0, columnspan=6, padx=10, pady=10)
        self.excel_aktar_button = ttk.Button(self.pencere, text="İstekleri Excel'e Aktar", command=self.talep_excel_aktar, style='TButton')
        self.excel_aktar_button.grid(row=10, column=0, columnspan=6, padx=10, pady=10)

        self.show_author_label(self.pencere)
        

    def ambar_sec(self, secim):
        self.ambar_var.set(secim)
        for widget in self.pencere.grid_slaves():
            if isinstance(widget, tk.Button) and widget.cget('text') in self.ambar_secimleri:
                if widget.cget('text') == secim:
                    widget.configure(bg='#66ff66')
                else:
                    widget.configure(bg='#f0f0f0')

    def talep_sec(self, secim):
        self.talep_var.set(secim)
        for widget in self.pencere.grid_slaves():
            if isinstance(widget, tk.Button) and widget.cget('text') in ["Evet", "Hayır"]:
                if widget.cget('text') == secim:
                    widget.configure(bg='#66ff66')
                else:
                    widget.configure(bg='#f0f0f0')

    def sevk_sec(self, secim):
        self.sevk_sekli_var.set(secim)
        for widget in self.pencere.grid_slaves():
            if isinstance(widget, tk.Button) and widget.cget('text') in ["Ring", "Kargo", "Otobüs"]:
                if widget.cget('text') == secim:
                    widget.configure(bg='#66ff66')
                else:
                    widget.configure(bg='#f0f0f0')

    def yerli_ithal_sec(self, secim):
        self.yerli_ithal_var.set(secim)
        for widget in self.pencere.grid_slaves():
            if isinstance(widget, tk.Button) and widget.cget('text') in ["Yerli", "İthal"]:
                if widget.cget('text') == secim:
                    widget.configure(bg='#66ff66')
                else:
                    widget.configure(bg='#f0f0f0')

    def urun_iste(self):
        ambar = self.ambar_var.get()
        kullanici_adi = self.kullanici_adi_entry.get()
        mrp = self.mrp_entry.get()
        miktar = self.miktar_entry.get()
        talep_edildi = self.talep_var.get()
        tarih = self.tarih_entry.get()
        sevk_sekli = self.sevk_sekli_var.get()
        yerli_ithal = self.yerli_ithal_var.get()

        talep_bilgileri = {
            'Ambar': ambar,
            'Ekip Kullanıcı Adı': kullanici_adi,
            'MRP': mrp,
            'Miktar': miktar,
            'Talep Edildi Mi': talep_edildi,
            'Talep Edilen Tarih': tarih,
            'Sevk Şekli': sevk_sekli,
            'Yerli mi İthal mi': yerli_ithal
        }
        self.urun_istekleri.append(talep_bilgileri)  # Listeye ekle
        self.talep_listbox.insert(tk.END, f"{kullanici_adi}: {miktar} adet {ambar} ambarından talep edildi.")
        messagebox.showinfo("Başarılı", "Ürün isteği başarıyla gönderildi.")

    def talep_excel_aktar(self):
        df = pd.DataFrame(self.urun_istekleri)
        dosya_yolu = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")])
        if dosya_yolu:
            df.to_excel(dosya_yolu, index=False)
            messagebox.showinfo("Başarılı", "İstekler başarıyla Excel dosyasına aktarıldı.")

    def urun_ekleme_penceresi(self):
        self.pencere = tk.Toplevel(self.root)
        self.pencere.title("Ürün Ekleme")
        self.pencere.geometry("400x300")
        self.pencere.configure(bg='#f0f0f0')

        self.urun_adi_label = tk.Label(self.pencere, text="Ürün Adı:", bg='#f0f0f0')
        self.urun_adi_label.pack()
        self.urun_adi_entry = tk.Entry(self.pencere)
        self.urun_adi_entry.pack()

        self.miktar_label = tk.Label(self.pencere, text="Miktar:", bg='#f0f0f0')
        self.miktar_label.pack()
        self.miktar_entry = tk.Entry(self.pencere)
        self.miktar_entry.pack()

        self.ekle_button = ttk.Button(self.pencere, text="Ürün Ekle", command=self.urun_ekle, style='TButton')
        self.ekle_button.pack()
        
    def urun_ekle(self):
        urun_adi = self.urun_adi_entry.get()
        miktar = int(self.miktar_entry.get())
        if urun_adi in self.depo:
            self.depo[urun_adi]['miktar'] += miktar
        else:
            self.depo[urun_adi] = {'miktar': miktar, 'lokasyonlar': []}
        messagebox.showinfo("Başarılı", f"{urun_adi} başarıyla eklendi. Mevcut miktar: {self.depo[urun_adi]['miktar']}")
        self.dosyayi_kaydet()
        self.stok_goster()

    def urun_cikarma_penceresi(self):
        self.pencere = tk.Toplevel(self.root)
        self.pencere.title("Ürün Çıkarma")
        self.pencere.geometry("400x400")
        self.pencere.configure(bg='#f0f0f0')

        self.urun_cikar_adi_label = tk.Label(self.pencere, text="Çıkarılacak Ürün Adı:", bg='#f0f0f0')
        self.urun_cikar_adi_label.pack()
        self.urun_cikar_adi_entry = tk.Entry(self.pencere)
        self.urun_cikar_adi_entry.pack()

        self.cikar_miktar_label = tk.Label(self.pencere, text="Miktar:", bg='#f0f0f0')
        self.cikar_miktar_label.pack()
        self.cikar_miktar_entry = tk.Entry(self.pencere)
        self.cikar_miktar_entry.pack()

        self.lokasyon_label = tk.Label(self.pencere, text="Gönderilen Lokasyon:", bg='#f0f0f0')
        self.lokasyon_label.pack()
        self.lokasyon_entry = tk.Entry(self.pencere)
        self.lokasyon_entry.pack()

        self.muhattap_label = tk.Label(self.pencere, text="Muhattap:", bg='#f0f0f0')
        self.muhattap_label.pack()
        self.muhattap_entry = tk.Entry(self.pencere)
        self.muhattap_entry.pack()

        self.cikar_button = ttk.Button(self.pencere, text="Ürün Çıkar", command=self.urun_cikar, style='TButton')
        self.cikar_button.pack()

    def urun_cikar(self):
        urun_adi = self.urun_cikar_adi_entry.get()
        miktar = int(self.cikar_miktar_entry.get())
        lokasyon = self.lokasyon_entry.get()
        muhattap = self.muhattap_entry.get()
        tarih_saat = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
        if urun_adi in self.depo:
            self.depo[urun_adi]['miktar'] -= miktar
            self.depo[urun_adi]['lokasyonlar'].append((lokasyon, miktar, muhattap, tarih_saat))
            if self.depo[urun_adi]['miktar'] <= 0:
                del self.depo[urun_adi]
            messagebox.showinfo("Başarılı", f"{urun_adi} başarıyla çıkarıldı. Kalan miktar: {self.depo.get(urun_adi, {'miktar': 0})['miktar']}")
        else:
            self.depo[urun_adi] = {'miktar': -miktar, 'lokasyonlar': [(lokasyon, miktar, muhattap, tarih_saat)]}
            messagebox.showinfo("Başarılı", f"{urun_adi} başarıyla çıkarıldı ve depo güncellendi. Mevcut miktar: {-miktar}")

        self.dosyayi_kaydet()
        self.stok_goster()


    def stok_goster(self):
        self.stok_listbox.delete(0, tk.END)
        for urun_adi, bilgiler in self.depo.items():
            self.stok_listbox.insert(tk.END, f"{urun_adi}: {bilgiler['miktar']} adet")
            for entry in bilgiler['lokasyonlar']:
                if len(entry) == 3:
                    lokasyon, miktar, tarih_saat = entry
                    self.stok_listbox.insert(tk.END, f"  - {miktar} adet {lokasyon} lokasyonuna gönderildi ({tarih_saat})")
                elif len(entry) == 4:
                    lokasyon, miktar, muhattap, tarih_saat = entry
                    self.stok_listbox.insert(tk.END, f"  - {miktar} adet {lokasyon} lokasyonuna gönderildi ({tarih_saat}), Muhattap: {muhattap}")

    def dosyayi_kaydet(self):
        with open('depo_verileri.json', 'w') as dosya:
            json.dump(self.depo, dosya)

    def dosyayi_yukle(self):
        if os.path.exists('depo_verileri.json'):
            with open('depo_verileri.json', 'r') as dosya:
                self.depo = json.load(dosya)

    def depo_sifirla_uyari(self):
        result = messagebox.askyesno("Depo Sıfırlama", "Depoyu kaydettiğinizden emin misiniz? Sıfırlama işlemi geri alınamaz.")
        if result:
            self.depo_sifirla()

    def depo_sifirla(self):
        self.depo = {}
        self.dosyayi_kaydet()
        self.stok_goster()
        messagebox.showinfo("Başarılı", "Depo başarıyla sıfırlandı.")

    def excel_aktar(self):
        data = []
        for urun_adi, bilgiler in self.depo.items():
            for entry in bilgiler['lokasyonlar']:
                if len(entry) == 4:
                    lokasyon, miktar, muhattap, tarih_saat = entry
                    data.append([urun_adi, miktar, lokasyon, muhattap, tarih_saat])
                elif len(entry) == 3:
                    lokasyon, miktar, tarih_saat = entry
                    data.append([urun_adi, miktar, lokasyon, "Bilinmiyor", tarih_saat])
                else:
                    messagebox.showerror("Hata", f"Beklenmeyen veri formatı: {entry}")
                    return

        df = pd.DataFrame(data, columns=["Ürün Adı", "Miktar", "Lokasyon", "Muhattap", "Tarih Saat"])
        dosya_yolu = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")])
        if dosya_yolu:
            df.to_excel(dosya_yolu, index=False)
            messagebox.showinfo("Başarılı", "İşlemler başarıyla Excel dosyasına aktarıldı.")

root = tk.Tk()
app = DepoYonetimi(root)
root.mainloop()
